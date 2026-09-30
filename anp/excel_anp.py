from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

import pandas as pd

from anp.base_anp import MESES_INV

FUENTE = Font(name="Calibri", size=11)
FUENTE_TITULO = Font(name="Calibri", size=11, bold=True)
CENTRO = Alignment(horizontal="center", vertical="center")
FORMATO_MONTO = "#,##0.00"


def _valor_celda(valor):
    if valor is None:
        return None
    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(valor, float) and valor == 0:
        return None
    return valor


def _titulo_mes(anio, mes, anios_multiples):
    nombre = MESES_INV[mes]
    if anios_multiples:
        return f"{nombre} {anio}"
    return nombre


def _escribir_resumen(ws, resumen):
    meses = resumen["meses"]
    anios_multiples = resumen["anios_multiples"]
    ws["A2"] = "Promotor"
    ws["A2"].font = FUENTE_TITULO
    columna = 2
    for anio, mes in meses:
        ws.merge_cells(
            start_row=1, start_column=columna, end_row=1, end_column=columna + 2
        )
        titulo = ws.cell(row=1, column=columna, value=_titulo_mes(anio, mes, anios_multiples))
        titulo.font = FUENTE_TITULO
        titulo.alignment = CENTRO
        for offset, texto in enumerate(("GMM", "VIDA", "TOTAL")):
            celda = ws.cell(row=2, column=columna + offset, value=texto)
            celda.font = FUENTE_TITULO
            celda.alignment = CENTRO
        columna += 3
    total_col = columna
    ws.merge_cells(start_row=1, start_column=total_col, end_row=2, end_column=total_col)
    total = ws.cell(row=1, column=total_col, value="TOTAL GENERAL")
    total.font = FUENTE_TITULO
    total.alignment = CENTRO

    for offset, fila in enumerate(resumen["filas"]):
        excel_fila = 3 + offset
        ws.cell(row=excel_fila, column=1, value=fila["promotor"]).font = FUENTE
        col = 2
        for anio, mes in meses:
            montos = fila["valores"].get((anio, mes), {})
            for campo in ("GMM", "Vida", "TOTAL"):
                celda = ws.cell(
                    row=excel_fila,
                    column=col,
                    value=_valor_celda(montos.get(campo)),
                )
                celda.font = FUENTE
                celda.number_format = FORMATO_MONTO
                col += 1
        general = ws.cell(
            row=excel_fila,
            column=total_col,
            value=_valor_celda(fila["total_general"]),
        )
        general.font = FUENTE
        general.number_format = FORMATO_MONTO

    ws.column_dimensions["A"].width = 14
    for indice in range(2, total_col + 1):
        ws.column_dimensions[get_column_letter(indice)].width = 12


def _valor_detalle(valor):
    if valor is None:
        return None
    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(valor, "to_pydatetime"):
        try:
            return valor.to_pydatetime()
        except (TypeError, ValueError):
            return valor
    return valor


def _escribir_detalle(ws, df):
    for r_idx, fila in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
        for c_idx, valor in enumerate(fila, start=1):
            celda = ws.cell(row=r_idx, column=c_idx, value=_valor_detalle(valor))
            celda.font = FUENTE_TITULO if r_idx == 1 else FUENTE


def guardar_base_anp(detalle, resumen, ruta):
    wb = Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    _escribir_resumen(ws_resumen, resumen)
    ws_detalle = wb.create_sheet("Detalle")
    _escribir_detalle(ws_detalle, detalle)
    wb.save(ruta)
    return str(ruta)
