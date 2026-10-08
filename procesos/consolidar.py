from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook


def _copiar_hoja(origen, destino, titulo=None):
    nombre = titulo or origen.title
    existente = {hoja.title for hoja in destino.worksheets}
    base = nombre
    indice = 2
    while nombre in existente:
        nombre = f"{base} ({indice})"
        indice += 1
    nueva = destino.create_sheet(nombre)
    for fila in origen.iter_rows():
        for celda in fila:
            nueva_celda = nueva.cell(
                row=celda.row, column=celda.column, value=celda.value
            )
            try:
                if getattr(celda, "has_style", False):
                    nueva_celda.font = copy(celda.font)
                    nueva_celda.fill = copy(celda.fill)
                    nueva_celda.border = copy(celda.border)
                    nueva_celda.alignment = copy(celda.alignment)
                    nueva_celda.number_format = celda.number_format
            except Exception:
                pass
    for rango in origen.merged_cells.ranges:
        nueva.merge_cells(str(rango))
    for letra, dim in origen.column_dimensions.items():
        if dim.width:
            nueva.column_dimensions[letra].width = dim.width
    return nueva


def consolidar_excels(rutas, ruta_salida):
    wb = Workbook()
    wb.remove(wb.active)
    for ruta in rutas:
        archivo = Path(ruta)
        if not archivo.exists():
            raise ValueError(f"No se encontró el archivo a consolidar: {archivo.name}")
        origen = load_workbook(archivo)
        for hoja in origen.worksheets:
            _copiar_hoja(hoja, wb)
    if not wb.worksheets:
        raise ValueError("No hay hojas para consolidar.")
    Path(ruta_salida).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_salida)
    return str(ruta_salida)
