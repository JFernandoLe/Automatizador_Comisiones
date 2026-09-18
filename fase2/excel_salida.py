from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils.dataframe import dataframe_to_rows

import pandas as pd

from fase2.transformaciones import (
    COLUMNAS_GMM,
    COLUMNAS_VIDA,
    construir_tabla_principal_bonos,
    construir_tablas_seleccionadas_bonos,
)

FUENTE = Font(name="Calibri", size=12)
FUENTE_TITULO = Font(name="Calibri", size=12, bold=True)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
FORMATO_MONTO = "#,##0.00"


def _aplicar_fuente(ws, celda, negrita=False):
    c = ws[celda]
    c.font = FUENTE_TITULO if negrita else FUENTE
    c.alignment = CENTRO


def _escribir_encabezados(ws):
    ws.merge_cells("A1:U1")
    ws["A1"] = "VIDA"
    _aplicar_fuente(ws, "A1", negrita=True)

    ws.merge_cells("D2:M2")
    ws["D2"] = "AGENTE"
    _aplicar_fuente(ws, "D2", negrita=True)

    ws.merge_cells("N2:T2")
    ws["N2"] = "PROMOTOR"
    _aplicar_fuente(ws, "N2", negrita=True)

    ws.merge_cells("D3:F3")
    ws["D3"] = "1ER. AÑO"
    ws.merge_cells("G3:I3")
    ws["G3"] = "OTROS CONCEPTOS"
    ws.merge_cells("J3:L3")
    ws["J3"] = "RENOVACIÓN"
    ws.merge_cells("M3:M4")
    ws["M3"] = "TOTAL"
    ws.merge_cells("N3:P3")
    ws["N3"] = "1ER. AÑO"
    ws.merge_cells("Q3:S3")
    ws["Q3"] = "RENOVACIÓN"
    ws.merge_cells("T3:T4")
    ws["T3"] = "TOTAL"
    ws.merge_cells("U3:U4")
    ws["U3"] = "TOTAL GENERAL"

    for celda in ("D3", "G3", "J3", "M3", "N3", "Q3", "T3", "U3"):
        _aplicar_fuente(ws, celda, negrita=True)

    vida_fila4 = [
        "prom",
        "Promotor",
        "Clave con PF",
        "META",
        "VIDA",
        "TOTAL",
        "META",
        "VIDA",
        "TOTAL",
        "META",
        "VIDA",
        "TOTAL",
        None,
        "META",
        "VIDA",
        "TOTAL",
        "META",
        "VIDA",
        "TOTAL",
        None,
        None,
    ]
    for indice, valor in enumerate(vida_fila4, start=1):
        if valor is None:
            continue
        celda = ws.cell(row=4, column=indice, value=valor)
        celda.font = FUENTE
        celda.alignment = CENTRO

    ws.merge_cells("W1:AH1")
    ws["W1"] = "GMM"
    _aplicar_fuente(ws, "W1", negrita=True)

    ws.merge_cells("Z3:AC3")
    ws["Z3"] = "AGENTE"
    ws.merge_cells("AD3:AG3")
    ws["AD3"] = "PROMOTOR"
    ws.merge_cells("AH3:AH4")
    ws["AH3"] = "TOTAL GENERAL"
    for celda in ("Z3", "AD3", "AH3"):
        _aplicar_fuente(ws, celda, negrita=True)

    gmm_fila4 = [
        "prom",
        "Promotor",
        "Clave con PF a PM",
        "1ER. AÑO",
        "OTROS CONCEPTOS",
        "RENOVACIÓN",
        "TOTAL",
        "1ER. AÑO",
        "OTROS CONCEPTOS",
        "RENOVACIÓN",
        "TOTAL",
        None,
    ]
    for indice, valor in enumerate(gmm_fila4, start=23):
        if valor is None:
            continue
        celda = ws.cell(row=4, column=indice, value=valor)
        celda.font = FUENTE
        celda.alignment = CENTRO


def _escribir_bloque(ws, df, columnas, col_inicio, filas_monto):
    for offset, fila in enumerate(dataframe_to_rows(df[columnas], index=False, header=False)):
        fila_excel = 5 + offset
        for indice, valor in enumerate(fila):
            celda = ws.cell(row=fila_excel, column=col_inicio + indice, value=valor)
            celda.font = FUENTE
            if indice in filas_monto:
                celda.number_format = FORMATO_MONTO


def _valor_celda(valor):
    if valor is None:
        return None
    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass
    return valor


def _nombre_columna_bonos(nombre):
    if str(nombre).strip() == "Ramo_clasif":
        return "Ramo"
    return str(nombre)


def _escribir_tabla_principal(ws, filas, col_inicio):
    encabezados = ("", "Vida", "GMM", "TOTAL GENERAL")
    for indice, valor in enumerate(encabezados):
        celda = ws.cell(row=1, column=col_inicio + indice, value=valor or None)
        celda.font = FUENTE_TITULO
        celda.alignment = CENTRO
    for offset, fila in enumerate(filas, start=2):
        etiqueta = ws.cell(
            row=offset, column=col_inicio, value=fila["etiqueta"]
        )
        etiqueta.font = FUENTE_TITULO if fila["seccion"] else FUENTE
        for extra, campo in enumerate(("vida", "gmm", "total"), start=1):
            celda = ws.cell(
                row=offset,
                column=col_inicio + extra,
                value=fila[campo],
            )
            celda.font = FUENTE_TITULO if fila["seccion"] else FUENTE
            celda.number_format = FORMATO_MONTO


SEPARACION_TABLAS = 3
ANCHO_TABLA_MINI = 2


def _escribir_tabla_combinacion(ws, tabla, col_inicio):
    ws.cell(row=1, column=col_inicio, value=tabla["figura"]).font = FUENTE_TITULO
    ws.cell(row=2, column=col_inicio, value=tabla["ramo2"]).font = FUENTE_TITULO
    ws.cell(row=3, column=col_inicio, value=tabla["ini_ren"]).font = FUENTE_TITULO
    encabezado_clave = ws.cell(
        row=4, column=col_inicio, value=tabla["nombre_clave"]
    )
    encabezado_importe = ws.cell(
        row=4, column=col_inicio + 1, value="Suma de importe"
    )
    encabezado_clave.font = FUENTE_TITULO
    encabezado_importe.font = FUENTE_TITULO
    encabezado_clave.alignment = CENTRO
    encabezado_importe.alignment = CENTRO
    for offset, fila in enumerate(tabla["filas"], start=5):
        ws.cell(
            row=offset, column=col_inicio, value=_valor_celda(fila["clave"])
        ).font = FUENTE
        celda_importe = ws.cell(
            row=offset,
            column=col_inicio + 1,
            value=_valor_celda(fila["importe"]),
        )
        celda_importe.font = FUENTE
        celda_importe.number_format = FORMATO_MONTO


def _escribir_bloque_combinaciones(ws, tablas, col_inicio):
    columna = col_inicio
    for tabla in tablas:
        _escribir_tabla_combinacion(ws, tabla, columna)
        columna += ANCHO_TABLA_MINI + SEPARACION_TABLAS
    return columna


def _escribir_pagos_bonos(wb, df, combinaciones=None):
    ws = wb.create_sheet("Pagos_de_Bonos")
    for indice, nombre in enumerate(df.columns, start=1):
        celda = ws.cell(row=1, column=indice, value=_nombre_columna_bonos(nombre))
        celda.font = FUENTE_TITULO
        celda.alignment = CENTRO
    for offset, fila in enumerate(
        dataframe_to_rows(df, index=False, header=False), start=2
    ):
        for indice, valor in enumerate(fila, start=1):
            celda = ws.cell(row=offset, column=indice, value=_valor_celda(valor))
            celda.font = FUENTE
            if "FECHA" in str(df.columns[indice - 1]).strip().upper():
                celda.number_format = "dd/mm/yyyy"

    col_resumen = df.shape[1] + 3
    _escribir_tabla_principal(
        ws, construir_tabla_principal_bonos(df), col_resumen
    )
    col_promotor = col_resumen + 4 + SEPARACION_TABLAS
    seleccion = combinaciones or {}
    _escribir_bloque_combinaciones(
        ws,
        construir_tablas_seleccionadas_bonos(
            df,
            figuras=seleccion.get("figuras"),
            ramos=seleccion.get("ramos"),
            conceptos=seleccion.get("conceptos"),
        ),
        col_promotor,
    )


def guardar_comision(
    df_vida, df_gmm, df_bonos, ruta="Comision.xlsx", combinaciones=None
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Comision"
    _escribir_encabezados(ws)
    _escribir_bloque(ws, df_vida, COLUMNAS_VIDA, 1, set(range(3, 21)))
    _escribir_bloque(ws, df_gmm, COLUMNAS_GMM, 23, set(range(3, 12)))

    anchos = {
        "A": 12,
        "B": 14,
        "C": 16,
        "U": 14,
        "W": 12,
        "X": 14,
        "Y": 18,
        "AA": 18,
        "AE": 18,
        "AH": 14,
    }
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    _escribir_pagos_bonos(wb, df_bonos, combinaciones=combinaciones)
    wb.save(ruta)
    return ruta
