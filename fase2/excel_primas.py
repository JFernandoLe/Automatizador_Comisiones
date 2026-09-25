from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import pandas as pd

FUENTE = Font(name="Calibri", size=12)
FUENTE_TITULO = Font(name="Calibri", size=12, bold=True)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
FORMATO_MONTO = "#,##0.00"
SEPARACION_TABLAS = 1


def _valor_celda(valor):
    if valor is None:
        return None
    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass
    return valor


def _ancho_tabla(tabla):
    return 4 + tabla["mes_hasta"]


def _escribir_tabla(ws, tabla, col_inicio):
    anio = ws.cell(row=1, column=col_inicio, value=tabla["anio"])
    titulo = ws.cell(row=1, column=col_inicio + 1, value=tabla["titulo"])
    anio.font = FUENTE_TITULO
    titulo.font = FUENTE_TITULO

    encabezados = ["Promotor", "Agente"] + list(range(1, tabla["mes_hasta"] + 1))
    encabezados.append("Total general")
    for indice, valor in enumerate(encabezados):
        celda = ws.cell(row=2, column=col_inicio + 1 + indice, value=valor)
        celda.font = FUENTE_TITULO
        celda.alignment = CENTRO

    for offset, registro in enumerate(tabla["filas"]):
        fila = 3 + offset
        ws.cell(
            row=fila,
            column=col_inicio + 1,
            value=_valor_celda(registro["promotor"]),
        ).font = FUENTE
        ws.cell(
            row=fila,
            column=col_inicio + 2,
            value=_valor_celda(registro["agente"]),
        ).font = FUENTE
        for mes in range(1, tabla["mes_hasta"] + 1):
            celda = ws.cell(
                row=fila,
                column=col_inicio + 2 + mes,
                value=_valor_celda(registro["meses"].get(mes) or None),
            )
            celda.font = FUENTE
            celda.number_format = FORMATO_MONTO
        total = ws.cell(
            row=fila,
            column=col_inicio + 3 + tabla["mes_hasta"],
            value=_valor_celda(registro["total"]),
        )
        total.font = FUENTE
        total.number_format = FORMATO_MONTO

    anchos = [10, 14, 14] + [14] * tabla["mes_hasta"] + [16]
    for indice, ancho in enumerate(anchos):
        letra = get_column_letter(col_inicio + indice)
        ws.column_dimensions[letra].width = ancho
    return col_inicio + _ancho_tabla(tabla) + SEPARACION_TABLAS


def guardar_primas(tablas, ruta="Primas.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bases PP"
    columna = 1
    for tabla in tablas:
        columna = _escribir_tabla(ws, tabla, columna)
    wb.save(ruta)
    return ruta
